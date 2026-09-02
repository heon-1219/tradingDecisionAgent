#!/usr/bin/env python
"""Convert a PDF / PPTX / text file into a markdown document for data/external/docs/.

Usage:
  python scripts/parse_docs.py <input.pdf|.pptx|.txt|.md> <output.md>
         [--title T] [--source URL] [--as-of YYYY-MM-DD] [--products "code1,code2"]

The output starts with the YAML front-matter that finagent.kb.docs expects:
  title / source / as_of / products.  PDFs get one "## Page N" section per page.
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

FRONT_MATTER_RE = re.compile(r"\A---\s*\n(.*?)\n---\s*\n?", re.DOTALL)


def pdf_to_md(path: Path) -> str:
    try:
        import pymupdf as fitz  # PyMuPDF >= 1.24
    except ImportError:  # pragma: no cover - older PyMuPDF
        import fitz
    parts = []
    with fitz.open(path) as doc:
        for n, page in enumerate(doc, start=1):
            parts.append(f"## Page {n}\n\n{page.get_text('text')}")
    return "\n\n".join(parts)


def pptx_to_md(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        sys.exit("error: .pptx input needs the optional 'python-pptx' package "
                 "(pip install python-pptx)")
    parts = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
        parts.append(f"## Slide {n}\n\n" + "\n".join(t for t in texts if t.strip()))
    return "\n\n".join(parts)


def text_to_md(path: Path) -> tuple[str, dict[str, str]]:
    """Pass-through; an existing front-matter block is stripped and its keys returned."""
    text = path.read_text(encoding="utf-8")
    meta: dict[str, str] = {}
    if m := FRONT_MATTER_RE.match(text):
        for line in m.group(1).splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip()] = v.strip()
        text = text[m.end():]
    return text, meta


def clean(text: str) -> str:
    text = "\n".join(line.rstrip() for line in text.splitlines())
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def front_matter(title: str, source: str, as_of: str, products: list[str]) -> str:
    codes = ", ".join(f'"{p}"' for p in products)
    return f"---\ntitle: {title}\nsource: {source}\nas_of: {as_of}\nproducts: [{codes}]\n---\n\n"


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", type=Path)
    ap.add_argument("output", type=Path)
    ap.add_argument("--title")
    ap.add_argument("--source")
    ap.add_argument("--as-of", dest="as_of")
    ap.add_argument("--products", default="", help='comma-separated product codes/names')
    a = ap.parse_args(argv)

    suffix = a.input.suffix.lower()
    meta: dict[str, str] = {}
    if suffix == ".pdf":
        body = pdf_to_md(a.input)
    elif suffix == ".pptx":
        body = pptx_to_md(a.input)
    elif suffix in {".txt", ".md", ".markdown"}:
        body, meta = text_to_md(a.input)
    else:
        sys.exit(f"error: unsupported input type '{suffix}' (use .pdf, .pptx, .txt or .md)")

    title = a.title or meta.get("title") or a.input.stem
    source = a.source or meta.get("source") or a.input.as_posix()
    as_of = a.as_of or meta.get("as_of") or dt.date.today().isoformat()
    products = [p.strip() for p in a.products.split(",") if p.strip()]
    if not products and (raw := meta.get("products", "")):
        products = [p.strip().strip('"\'') for p in raw.strip("[]").split(",") if p.strip()]

    a.output.parent.mkdir(parents=True, exist_ok=True)
    a.output.write_text(front_matter(title, source, as_of, products) + clean(body), encoding="utf-8")
    print(f"wrote {a.output} ({a.output.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
